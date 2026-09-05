"""M2-3 自查：回溯还原。"""
from __future__ import annotations

import pytest
from docx import Document

from app.core.mapping.crypto import export_encrypted
from app.core.mapping.exporter import build_mapping_json
from app.core.mapping.map_table import MapTable
from app.core.mapping.restorer import RestoreError, restore_document


@pytest.fixture
def batch(tmp_path):
    """脱敏导出 + 加密映射，返回 (脱敏文档路径, 映射路径, 输出目录)。"""
    src = tmp_path / "合同.docx"
    doc = Document()
    doc.add_paragraph("出借人：张三，借款人：李四。")
    doc.save(str(src))

    table = MapTable()
    table.add_finding("张三", "person", "lexicon", 0.85, pending=False, file_name="合同.docx")
    table.add_finding("李四", "person", "lexicon", 0.85, pending=False, file_name="合同.docx")

    from app.core.redact.docx_redactor import redact_document
    redacted = tmp_path / "合同_脱敏.docx"
    redact_document(str(src), str(redacted), table.effective_pairs("合同.docx"))

    map_path = tmp_path / "映射.json"
    export_encrypted(build_mapping_json(table, [str(src)]), "pw123", map_path)
    return str(redacted), str(map_path), str(tmp_path / "还原输出")


def test_restore_roundtrip(batch):
    """脱敏 → 还原 → 与原文一致。"""
    redacted, map_path, out_dir = batch
    result = restore_document(redacted, map_path, "pw123", out_dir)
    from app.core.parser.docx_parser import extract_paragraphs
    text = "\n".join(extract_paragraphs(result.output_path))
    assert "张三" in text and "李四" in text
    assert "甲某" not in text and "乙某" not in text
    assert result.replaced >= 2 and not result.unrestored


def test_restore_after_ai_edit(batch):
    """AI 修改（增删段落、保留代称）后仍可还原；被删代称进未还原项。"""
    redacted, map_path, out_dir = batch
    doc = Document(redacted)
    doc.add_paragraph("AI 新增段落：甲某承诺按期还款。")
    # 删除含"乙某"的句子（模拟 AI 删除）
    for p in doc.paragraphs:
        if "乙某" in p.text:
            for r in p.runs:
                r.text = r.text.replace("，借款人：乙某", "")
    doc.save(redacted)

    result = restore_document(redacted, map_path, "pw123", out_dir)
    from app.core.parser.docx_parser import extract_paragraphs
    text = "\n".join(extract_paragraphs(result.output_path))
    assert "张三" in text                       # 保留的代称已还原
    aliases = [u["alias"] for u in result.unrestored]
    assert "乙某" in aliases                    # 被删的代称进了未还原项


def test_wrong_password_rejected(batch):
    redacted, map_path, out_dir = batch
    with pytest.raises(RestoreError):
        restore_document(redacted, map_path, "bad", out_dir)


def test_corrupted_mapping_rejected(batch, tmp_path):
    bad = tmp_path / "bad.json"
    bad.write_text("not json", encoding="utf-8")
    with pytest.raises(RestoreError):
        restore_document(batch[0], str(bad), "pw123", batch[2])


def test_unknown_doc_name_warns(batch, tmp_path):
    redacted, map_path, out_dir = batch
    renamed = tmp_path / "改名.docx"
    import shutil
    shutil.copy(redacted, renamed)
    result = restore_document(str(renamed), map_path, "pw123", out_dir)
    assert result.name_warning


def test_txt_restore(tmp_path):
    """txt 直接字符串替换。"""
    doc = tmp_path / "ai产出.txt"
    doc.write_text("甲某与乙某签订补充协议。", encoding="utf-8")
    map_path = tmp_path / "m.json"
    export_encrypted({
        "version": "1.0", "batch_id": "b", "doc_manifest": [],
        "mappings": [
            {"id": 1, "category": "person", "original": "张三", "alias": "甲某",
             "occurrences": 1, "enabled": True},
            {"id": 2, "category": "person", "original": "李四", "alias": "乙某",
             "occurrences": 1, "enabled": True},
        ],
    }, "pw", map_path)
    result = restore_document(str(doc), str(map_path), "pw", str(tmp_path / "out"))
    from pathlib import Path
    text = Path(result.output_path).read_text(encoding="utf-8")
    assert text == "张三与李四签订补充协议。"
